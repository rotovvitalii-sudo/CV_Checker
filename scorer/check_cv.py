#!/usr/bin/env python3
"""
check_cv.py — script part of the CV checker.

1. Parses the .docx the way an ATS would (python-docx, paragraphs only) and verifies it is readable.
2. Checks ATS structure: no tables/images, standard section headings present.
3. Estimates length (one-pager budget).
4. Computes keyword coverage against a keywords JSON:
     {"must": ["subscription model", "P&L|profit and loss", ...],
      "nice": ["EdTech", ...]}
   A keyword may list alternatives separated by "|". Matching is case-insensitive,
   whole-word, tolerant to simple plural/-ing/-ed suffixes.

5. Renders the .docx to PDF with LibreOffice (soffice --headless) and counts pages — the hard
   one-page check. Also reports free space at the bottom of the last page (safety margin for
   renderers with different font metrics: Word / Google Docs / LibreOffice).

Usage:
  check_cv.py --docx CV.docx --keywords 01-keywords.json [--report 03-score.md] [--json] [--no-pdf]

Exit code 0 always; the pipeline reads the numbers, not the exit code.
"""
import argparse, json, math, re, shutil, subprocess, sys
from pathlib import Path

try:
    import docx  # python-docx
except ImportError:
    sys.exit("python-docx missing: run  .venv/bin/pip install python-docx")

STANDARD_HEADINGS = ["experience", "education"]
ONE_PAGE_WORDS = 620          # soft budget (master CV ≈ 560 words)
ONE_PAGE_LINES = 66           # A4, 10pt, 0.5" top/bottom margins, build_docx.js template (estimate only)
CHARS_PER_LINE = 100          # ~7.2" text width at 10pt Arial (estimate only; PDF render is the real check)
MIN_BOTTOM_FREE_MM = 12       # keep ≥ this much free at the bottom so other renderers don't spill to page 2
MUST_WEIGHT, NICE_WEIGHT = 2, 1


def render_pdf(docx_path: Path):
    """docx → pdf via LibreOffice. Returns (pdf_path, pages, bottom_free_mm) or (None, None, None)."""
    soffice = shutil.which("soffice")
    if not soffice:
        return None, None, None
    outdir = docx_path.parent
    subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", str(outdir), str(docx_path)],
                   capture_output=True, text=True, timeout=180)
    pdf = outdir / (docx_path.stem + ".pdf")
    if not pdf.exists():
        return None, None, None
    from pypdf import PdfReader
    r = PdfReader(str(pdf))
    pages = len(r.pages)
    last = r.pages[-1]
    height_pt = float(last.mediabox.height)
    lowest = [height_pt]

    def visitor(text, cm, tm, font_dict, font_size):
        if text.strip():
            lowest[0] = min(lowest[0], tm[5])

    last.extract_text(visitor_text=visitor)
    bottom_free_mm = round(lowest[0] * 25.4 / 72, 1)   # distance from page bottom to lowest text baseline
    return pdf, pages, bottom_free_mm


def extract(docx_path: Path):
    d = docx.Document(str(docx_path))
    paras = [p.text.strip() for p in d.paragraphs if p.text.strip()]
    tables = len(d.tables)
    images = len(d.inline_shapes)
    headings = []
    est_lines = 0.0
    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        # python-docx may not resolve styles written by other generators; read pStyle from XML directly
        ppr = p._p.pPr
        ps = ppr.pStyle if ppr is not None else None
        style_id = (ps.val if ps is not None else (p.style.name if p.style is not None else "")) or ""
        is_h1 = style_id.lower().replace(" ", "") == "heading1"
        if is_h1:
            headings.append(t)
        is_bullet = ppr is not None and ppr.numPr is not None
        est_lines += math.ceil(len(t) / (CHARS_PER_LINE - (6 if is_bullet else 0)))
        if is_h1:
            est_lines += 0.6  # heading spacing + border
    return paras, tables, images, headings, est_lines


def term_regex(term: str) -> re.Pattern:
    alts = []
    for alt in term.split("|"):
        alt = alt.strip()
        if not alt:
            continue
        words = [re.escape(w) for w in alt.split()]
        # last word tolerant to s / es / ing / ed
        words[-1] = words[-1] + r"(?:s|es|ing|ed)?"
        alts.append(r"\b" + r"[\s\-/]+".join(words) + r"\b")
    return re.compile("|".join(alts), re.IGNORECASE)


def coverage(text: str, terms: list[str]):
    hit, miss = [], []
    for t in terms:
        (hit if term_regex(t).search(text) else miss).append(t)
    return hit, miss


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--docx", required=True)
    ap.add_argument("--keywords", required=True)
    ap.add_argument("--report")
    ap.add_argument("--json", action="store_true")
    ap.add_argument("--no-pdf", action="store_true", help="skip LibreOffice render")
    a = ap.parse_args()

    paras, tables, images, headings, est_lines = extract(Path(a.docx))
    text = "\n".join(paras)
    words = len(text.split())

    pdf, pages, bottom_free = (None, None, None) if a.no_pdf else render_pdf(Path(a.docx))
    if pages is not None:
        length_ok = pages == 1 and bottom_free >= MIN_BOTTOM_FREE_MM
        length_note = (f"{pages} page(s), {bottom_free} mm free at bottom (min {MIN_BOTTOM_FREE_MM})"
                       if pages == 1 else f"{pages} pages — spills over")
    else:
        length_ok = round(est_lines) <= ONE_PAGE_LINES
        length_note = f"no PDF render (soffice missing) — estimate ~{est_lines:.0f}/{ONE_PAGE_LINES} lines"
    kw = json.loads(Path(a.keywords).read_text())
    must, nice = kw.get("must", []), kw.get("nice", [])

    must_hit, must_miss = coverage(text, must)
    nice_hit, nice_miss = coverage(text, nice)
    total_w = MUST_WEIGHT * len(must) + NICE_WEIGHT * len(nice)
    score = (MUST_WEIGHT * len(must_hit) + NICE_WEIGHT * len(nice_hit)) / total_w if total_w else 0

    # W1: верхняя треть — всё до раздела Experience (шапка, Career Profile, Highlights). Рекрутер смотрит сюда первым.
    exp_i = next((i for i, t in enumerate(paras) if t.strip().lower() == "experience"), None)
    top_text = "\n".join(paras[:exp_i]) if exp_i else ""
    top_hit, _ = coverage(top_text, must)
    # Главная задача роли: термины из 01-keywords.json → "task" (пишет tailor-cv); для старых ключей — первые 3 must
    task_terms = kw.get("task") or must[:3]
    task_hit = [t for t in task_terms if term_regex(t).search(top_text)]
    need_top = 2 if len(must) < 8 else 3
    top_third = {"must_hit": top_hit, "must_hit_n": len(top_hit), "need": need_top,
                 "task_terms": task_terms, "task_visible": bool(task_hit),
                 "pass": bool(top_text) and len(top_hit) >= need_top and bool(task_hit)}

    hl = [h.lower() for h in headings]
    missing_headings = [h for h in STANDARD_HEADINGS if not any(h in x for x in hl)]

    result = {
        "parse_ok": len(paras) > 0,
        "paragraphs": len(paras),
        "words": words,
        "words_budget": ONE_PAGE_WORDS,
        "est_lines": round(est_lines, 1),
        "lines_budget": ONE_PAGE_LINES,
        "pdf": str(pdf) if pdf else None,
        "pdf_pages": pages,
        "bottom_free_mm": bottom_free,
        "length_ok": length_ok,
        "length_note": length_note,
        "tables": tables,
        "images": images,
        "headings": headings,
        "missing_standard_headings": missing_headings,
        "must": {"hit": must_hit, "miss": must_miss, "rate": round(len(must_hit) / len(must), 3) if must else None},
        "nice": {"hit": nice_hit, "miss": nice_miss, "rate": round(len(nice_hit) / len(nice), 3) if nice else None},
        "score": round(score, 3),
        "top_third": top_third,
    }

    if a.report:
        md = [f"# Score — {Path(a.docx).name}", ""]
        md.append(f"**Overall keyword coverage:** {result['score']*100:.0f}%  (must ×{MUST_WEIGHT}, nice ×{NICE_WEIGHT})")
        md.append(f"**Must-have:** {len(must_hit)}/{len(must)} · **Nice-to-have:** {len(nice_hit)}/{len(nice)}")
        md.append("")
        tt = result["top_third"]
        md.append(f"**Top third (before Experience):** {'OK' if tt['pass'] else 'FAIL'} — main task "
                  f"{'visible' if tt['task_visible'] else 'NOT visible'} ({', '.join(tt['task_terms'])}); "
                  f"must keywords {tt['must_hit_n']} (need {tt['need']})")
        md.append("")
        md.append("## ATS structure")
        md.append(f"- Parsed paragraphs: {len(paras)} — {'OK' if result['parse_ok'] else 'FAIL'}")
        md.append(f"- Length: {length_note}; {words} words — {'OK (one page)' if length_ok else 'FAIL'}")
        md.append(f"- Tables: {tables}, images: {images} — {'OK' if tables == 0 and images == 0 else 'FAIL'}")
        md.append(f"- Headings: {', '.join(headings) or '—'}")
        if missing_headings:
            md.append(f"- Missing standard headings: {', '.join(missing_headings)}")
        md.append("")
        md.append("## Must-have — missing")
        md += [f"- {t}" for t in must_miss] or ["- none"]
        md.append("")
        md.append("## Nice-to-have — missing")
        md += [f"- {t}" for t in nice_miss] or ["- none"]
        md.append("")
        md.append("## Matched")
        md.append("- must: " + (", ".join(must_hit) or "—"))
        md.append("- nice: " + (", ".join(nice_hit) or "—"))
        Path(a.report).write_text("\n".join(md) + "\n")

    if a.json:
        print(json.dumps(result, ensure_ascii=False, indent=2))
    else:
        print(f"score={result['score']:.2f} must={len(must_hit)}/{len(must)} nice={len(nice_hit)}/{len(nice)} "
              f"pages={pages if pages is not None else '?'} bottom_free={bottom_free if bottom_free is not None else '?'}mm "
              f"words={words} tables={tables} images={images} "
              f"parse={'ok' if result['parse_ok'] else 'FAIL'} length={'ok' if length_ok else 'FAIL'}")


if __name__ == "__main__":
    main()
