#!/usr/bin/env python3
"""
cv_text.py — извлекает текст из загруженного CV (PDF или DOCX), без интерпретации.

Нужен, чтобы до запуска агента понять, есть ли в файле текст (скан PDF без текстового слоя
отсекается сразу), и чтобы агент работал с тем же текстом, что видит ATS.

Usage:
  cv_text.py --file CV.pdf [--out text.txt] [--json]
"""
import argparse, json
from pathlib import Path


def pdf_text(p: Path) -> str:
    from pypdf import PdfReader
    return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages)


def docx_text(p: Path) -> str:
    import docx
    d = docx.Document(str(p))
    lines = []
    for para in d.paragraphs:
        t = para.text.strip()
        if not t:
            continue
        style = (para.style.name if para.style is not None else "").lower()
        if style.startswith("heading") or style == "title":
            t = f"[{para.style.name}] {t}"
        elif "list" in style:
            t = f"- {t}"
        lines.append(t)
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    return "\n".join(lines)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--file", required=True)
    ap.add_argument("--out")
    ap.add_argument("--json", action="store_true")
    a = ap.parse_args()

    p = Path(a.file)
    ext = p.suffix.lower()
    try:
        text = pdf_text(p) if ext == ".pdf" else docx_text(p) if ext == ".docx" else ""
    except Exception as e:  # битый или зашифрованный файл
        print(json.dumps({"file": p.name, "ok": False, "words": 0,
                          "error": f"не удалось прочитать {ext.lstrip('.').upper()}: файл повреждён или защищён ({type(e).__name__})"},
                         ensure_ascii=False))
        return
    words = len(text.split())
    if a.out:
        Path(a.out).write_text(text)
    res = {"file": p.name, "format": ext.lstrip("."), "words": words, "ok": words >= 150,
           "error": None if words >= 150 else
           ("неподдерживаемый формат — нужен PDF или DOCX" if ext not in (".pdf", ".docx") else
            "в файле почти нет текста (скан?) — загрузи DOCX или PDF с текстовым слоем")}
    print(json.dumps(res, ensure_ascii=False) if a.json else f"{p.name}: {words} слов, ok={res['ok']}")


if __name__ == "__main__":
    main()
